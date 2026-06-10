"""Declarative XLSX workbook layout (engine-agnostic)."""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Protocol

from grid_view_spec.export.print import GridViewTablePrintContext, print_table_row_cells

if TYPE_CHECKING:
    import xlsxwriter
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from xlsxwriter.format import Format

XlsxCell = str | int | float | bool | None
XlsxRow = Sequence[XlsxCell]


@dataclass(frozen=True, slots=True)
class XlsxMergeRange:
    """Zero-based inclusive cell range (same as Excel / xlsxwriter)."""

    first_row: int
    first_col: int
    last_row: int
    last_col: int


@dataclass(frozen=True, slots=True)
class XlsxColWidth:
    """Column width in Excel character units."""

    col: int
    width: float


@dataclass(slots=True)
class XlsxSheet:
    """One worksheet: optional title banner, table, optional footer."""

    name: str
    title_rows: Sequence[XlsxRow] = field(default_factory=tuple)
    header_rows: Sequence[XlsxRow] = field(default_factory=tuple)
    data_rows: Sequence[XlsxRow] = field(default_factory=tuple)
    footer_rows: Sequence[XlsxRow] = field(default_factory=tuple)
    merges: list[XlsxMergeRange] = field(default_factory=list)
    col_widths: list[XlsxColWidth] = field(default_factory=list)
    freeze_panes: tuple[int, int] | None = None


@dataclass(slots=True)
class XlsxReport:
    """Workbook with one or more sheets."""

    sheets: list[XlsxSheet] = field(default_factory=list)


class XlsxBackend(Protocol):
    def render(self, report: XlsxReport) -> bytes: ...


def report_from_print_context(
    ctx: GridViewTablePrintContext,
    *,
    sheet_name: str = "Data",
    title_rows: Sequence[Sequence[XlsxCell]] = (),
    meta_lines: Sequence[str] = (),
) -> XlsxReport:
    """Convert a printable table context to a single-sheet workbook layout."""
    header_rows: list[XlsxRow] = [
        tuple(str(cell["label"]) for cell in header_row) for header_row in ctx["header_rows"]
    ]
    data_rows: list[XlsxRow] = [tuple(print_table_row_cells(row)) for row in ctx["rows"]]
    footer_rows: list[XlsxRow] = []
    footer_cells = ctx.get("footer_cells")
    if footer_cells:
        row: list[XlsxCell] = []
        for cell in footer_cells:
            colspan = int(cell.get("colspan", 1))
            row.append(cell["text"])
            for _ in range(max(0, colspan - 1)):
                row.append("")
        footer_rows.append(tuple(row))

    normalized_title_rows: tuple[XlsxRow, ...] = tuple(tuple(row) for row in title_rows)
    if meta_lines:
        merged_rows: list[XlsxRow] = list(normalized_title_rows)
        for line in meta_lines:
            merged_rows.append((line,))
        normalized_title_rows = tuple(merged_rows)

    sheet = XlsxSheet(
        name=sheet_name[:31],
        title_rows=normalized_title_rows,
        header_rows=tuple(header_rows),
        data_rows=tuple(data_rows),
        footer_rows=tuple(footer_rows),
    )
    return XlsxReport(sheets=[sheet])


def render_xlsx_workbook(report: XlsxReport, *, engine: str | None = None) -> bytes:
    """Render ``XlsxReport`` to bytes via configured backend."""
    from django.conf import settings

    name = engine or getattr(settings, "GRID_VIEW_XLSX_ENGINE", "xlsxwriter")
    backend = get_xlsx_backend(str(name))
    return backend.render(report)


def xlsx_response_from_report(
    report: XlsxReport,
    filename: str,
    *,
    engine: str | None = None,
):
    from django.http import HttpResponse

    xlsx_bytes = render_xlsx_workbook(report, engine=engine)
    response = HttpResponse(
        xlsx_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def get_xlsx_backend(name: str = "xlsxwriter") -> XlsxBackend:
    if name == "openpyxl":
        return OpenpyxlBackend()
    return XlsxWriterBackend()


class XlsxWriterBackend:
    """Write ``XlsxReport`` with xlsxwriter."""

    def render(self, report: XlsxReport) -> bytes:
        import xlsxwriter

        buffer = io.BytesIO()
        workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
        header_fmt = workbook.add_format({"bold": True, "bg_color": "#E8EEF4", "border": 1})
        title_fmt = workbook.add_format({"bold": True, "font_size": 12})
        footer_fmt = workbook.add_format({"bold": True, "top": 2})
        for sheet_def in report.sheets:
            self._write_sheet(workbook, sheet_def, header_fmt, title_fmt, footer_fmt)
        workbook.close()
        return buffer.getvalue()

    def _write_sheet(
        self,
        workbook: xlsxwriter.Workbook,
        sheet_def: XlsxSheet,
        header_fmt: Format,
        title_fmt: Format,
        footer_fmt: Format,
    ) -> None:
        name = (sheet_def.name or "Sheet1")[:31]
        ws = workbook.add_worksheet(name)
        row_idx = 0

        for title_row in sheet_def.title_rows:
            for col_idx, value in enumerate(title_row):
                ws.write(row_idx, col_idx, value, title_fmt)
            row_idx += 1
        if sheet_def.title_rows:
            row_idx += 1

        header_start = row_idx
        for header_row in sheet_def.header_rows:
            for col_idx, value in enumerate(header_row):
                ws.write(row_idx, col_idx, value, header_fmt)
            row_idx += 1

        for data_row in sheet_def.data_rows:
            for col_idx, value in enumerate(data_row):
                ws.write(row_idx, col_idx, value)
            row_idx += 1

        for footer_row in sheet_def.footer_rows:
            for col_idx, value in enumerate(footer_row):
                ws.write(row_idx, col_idx, value, footer_fmt)
            row_idx += 1

        for merge in sheet_def.merges:
            ws.merge_range(
                merge.first_row,
                merge.first_col,
                merge.last_row,
                merge.last_col,
                "",
            )

        for col_width in sheet_def.col_widths:
            ws.set_column(col_width.col, col_width.col, col_width.width)

        if sheet_def.freeze_panes is not None:
            first_row, first_col = sheet_def.freeze_panes
            ws.freeze_panes(first_row, first_col)
        elif sheet_def.header_rows:
            ws.freeze_panes(header_start + len(sheet_def.header_rows), 0)


class OpenpyxlBackend:
    """Write ``XlsxReport`` with openpyxl."""

    def render(self, report: XlsxReport) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        default_sheet = wb.active
        if default_sheet is not None:
            wb.remove(default_sheet)

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="E8EEF4")
        title_font = Font(bold=True, size=12)
        footer_font = Font(bold=True)

        for sheet_def in report.sheets:
            self._write_sheet(wb, sheet_def, header_font, header_fill, title_font, footer_font)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def _write_sheet(
        self,
        wb: Workbook,
        sheet_def: XlsxSheet,
        header_font: Font,
        header_fill: PatternFill,
        title_font: Font,
        footer_font: Font,
    ) -> None:
        from openpyxl.utils import get_column_letter

        name = (sheet_def.name or "Sheet1")[:31]
        ws = wb.create_sheet(title=name)
        row_idx = 1

        for title_row in sheet_def.title_rows:
            for col_idx, value in enumerate(title_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = title_font
            row_idx += 1
        if sheet_def.title_rows:
            row_idx += 1

        for header_row in sheet_def.header_rows:
            for col_idx, value in enumerate(header_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = header_font
                cell.fill = header_fill
            row_idx += 1

        for data_row in sheet_def.data_rows:
            for col_idx, value in enumerate(data_row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

        for footer_row in sheet_def.footer_rows:
            for col_idx, value in enumerate(footer_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = footer_font
            row_idx += 1

        for merge in sheet_def.merges:
            ws.merge_cells(
                start_row=merge.first_row + 1,
                start_column=merge.first_col + 1,
                end_row=merge.last_row + 1,
                end_column=merge.last_col + 1,
            )

        for col_width in sheet_def.col_widths:
            letter = get_column_letter(col_width.col + 1)
            ws.column_dimensions[letter].width = col_width.width

        if sheet_def.freeze_panes is not None:
            first_row, first_col = sheet_def.freeze_panes
            ws.freeze_panes = ws.cell(row=first_row + 1, column=first_col + 1)
