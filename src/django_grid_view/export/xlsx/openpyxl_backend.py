"""openpyxl backend — same declarative ``XlsxReport``; future template fill."""

from __future__ import annotations

import io
from pathlib import Path
from typing import TYPE_CHECKING

from django_grid_view.export.xlsx.layout import XlsxReport, XlsxSheet

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill


class OpenpyxlBackend:
    """Write ``XlsxReport`` with openpyxl.

    Template-based export (load branded ``.xlsx`` and fill named ranges) will live here
    when ``GRID_VIEW_XLSX_ENGINE=openpyxl`` and host apps pass a template path.
    """

    def render(self, report: XlsxReport) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        default_sheet = wb.active
        if default_sheet is not None:
            wb.remove(default_sheet)

        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="E8EEF4")
        title_font = Font(bold=True, size=12)
        footer_font = Font(bold=True)

        for sheet_def in report.sheets:
            self._write_sheet(wb, sheet_def, header_font, header_fill, title_font, footer_font)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def render_from_template(
        self,
        template_path: Path,
        report: XlsxReport,
        *,
        sheet_name: str | None = None,
    ) -> bytes:
        """Reserved for branded workbook templates (not implemented yet)."""
        msg = (
            "openpyxl template export is not implemented; use declarative XlsxReport "
            "or set GRID_VIEW_XLSX_ENGINE=xlsxwriter"
        )
        raise NotImplementedError(msg)

    def _write_sheet(
        self,
        wb: Workbook,
        sheet_def: XlsxSheet,
        header_font: Font,
        header_fill: PatternFill,
        title_font: Font,
        footer_font: Font,
    ) -> None:
        from openpyxl.utils import get_column_letter

        name = (sheet_def.name or "Sheet1")[:31]
        ws = wb.create_sheet(title=name)
        row_idx = 1

        for title_row in sheet_def.title_rows:
            for col_idx, value in enumerate(title_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = title_font
            row_idx += 1
        if sheet_def.title_rows:
            row_idx += 1

        header_start = row_idx
        for header_row in sheet_def.header_rows:
            for col_idx, value in enumerate(header_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = header_font
                cell.fill = header_fill
            row_idx += 1

        for data_row in sheet_def.data_rows:
            for col_idx, value in enumerate(data_row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

        for footer_row in sheet_def.footer_rows:
            for col_idx, value in enumerate(footer_row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = footer_font
            row_idx += 1

        for merge in sheet_def.merges:
            ws.merge_cells(
                start_row=merge.first_row + 1,
                start_column=merge.first_col + 1,
                end_row=merge.last_row + 1,
                end_column=merge.last_col + 1,
            )

        for col_width in sheet_def.col_widths:
            letter = get_column_letter(col_width.col + 1)
            ws.column_dimensions[letter].width = col_width.width

        if sheet_def.freeze_panes is not None:
            fr, fc = sheet_def.freeze_panes
            ws.freeze_panes = ws.cell(row=fr + 1, column=fc + 1)

        _ = header_start
